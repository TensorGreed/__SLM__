"""Generic intake (Wave 2a): any file a newcomer brings lands correctly.

Pins:
  * tabular_io reads CSV (cp1252 / BOM / ``;`` delimiters), TSV, JSON
    wrappers, JSON-that-is-JSONL, XLSX and Parquet into rows;
  * parsers raise instead of returning "[PDF parse error…]" strings — a
    scanned PDF / broken file ends up ERROR, never ACCEPTED training text;
    HTML drops scripts; DOCX keeps table cells;
  * a labelled CSV uploaded through Data → Upload and cleaned keeps its
    columns (one cleaned row per record, PII redacted per field, duplicates
    dropped) instead of being flattened and chunked at 1000 chars;
  * dataset prep never reads PDF/DOCX bytes as text lines;
  * the import wizard can take a browser upload (``file:`` locator) of any
    row format, and rejects documents with a pointer to the right place.
"""

from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient

from app.config import settings
from app.main import app
from app.services.dataset_service import _load_records_from_file
from app.utils.file_parsers import DocumentParseError, parse_file
from app.utils.tabular_io import load_rows

CP1252_CSV = "question;answer\nWhat does the café charge?;Five € per cup\nRefund window?;30 days\n".encode("cp1252")


def setUpModule():
    global _client_cm, client, _prev_auth
    _prev_auth = settings.AUTH_ENABLED
    settings.AUTH_ENABLED = False
    _client_cm = TestClient(app)
    client = _client_cm.__enter__()


def tearDownModule():
    _client_cm.__exit__(None, None, None)
    settings.AUTH_ENABLED = _prev_auth


def _blank_pdf() -> bytes:
    from PyPDF2 import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


class TabularIoTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _write(self, name: str, data: bytes | str) -> Path:
        path = self.root / name
        if isinstance(data, str):
            path.write_text(data, encoding="utf-8")
        else:
            path.write_bytes(data)
        return path

    def test_cp1252_semicolon_csv(self):
        rows = load_rows(self._write("excel.csv", CP1252_CSV))
        self.assertEqual(rows[0], {"question": "What does the café charge?", "answer": "Five € per cup"})
        self.assertEqual(len(rows), 2)

    def test_utf8_bom_and_blank_rows(self):
        rows = load_rows(self._write("bom.csv", "﻿text,label\nhi,pos\n,\n"))
        self.assertEqual(rows, [{"text": "hi", "label": "pos"}])

    def test_tsv(self):
        rows = load_rows(self._write("d.tsv", "text\tlabel\na, b\tneg\n"))
        self.assertEqual(rows, [{"text": "a, b", "label": "neg"}])

    def test_json_wrapper_and_jsonl_disguised_as_json(self):
        wrapped = load_rows(self._write("w.json", json.dumps({"data": [{"q": "1"}, {"q": "2"}]})))
        self.assertEqual([r["q"] for r in wrapped], ["1", "2"])
        lines = load_rows(self._write("l.json", '{"q": "a"}\n{"q": "b"}\n'))
        self.assertEqual([r["q"] for r in lines], ["a", "b"])

    def test_xlsx_and_parquet(self):
        import pandas as pd

        frame = pd.DataFrame({"question": ["Refund window?", "Hours?"], "answer": ["30 days", None], "n": [1, 2]})
        xlsx = self.root / "sheet.xlsx"
        frame.to_excel(xlsx, index=False)
        parquet = self.root / "data.parquet"
        frame.to_parquet(parquet)
        for path in (xlsx, parquet):
            with self.subTest(path.suffix):
                rows = load_rows(path)
                self.assertEqual(rows[0]["question"], "Refund window?")
                self.assertIsNone(rows[1]["answer"])
                self.assertEqual(rows[1]["n"], 2)
                json.dumps(rows)  # plain JSON values, no numpy types


class ParserTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_scanned_pdf_raises_readable_error(self):
        path = self.root / "scan.pdf"
        path.write_bytes(_blank_pdf())
        with self.assertRaisesRegex(DocumentParseError, "scanned"):
            parse_file(path)

    def test_broken_pdf_raises(self):
        path = self.root / "broken.pdf"
        path.write_bytes(b"not a pdf at all")
        with self.assertRaises(DocumentParseError):
            parse_file(path)

    def test_html_drops_scripts_and_keeps_text(self):
        path = self.root / "page.html"
        path.write_text(
            "<html><head><title>x</title><script>var secret=1;</script></head>"
            "<body><h1>Returns</h1><p>Refunds take&nbsp;30 days.</p><style>p{}</style></body></html>"
        )
        text = parse_file(path)
        self.assertIn("Returns", text)
        self.assertIn("Refunds take 30 days.", text)
        self.assertNotIn("secret", text)

    def test_docx_includes_tables(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("Warranty policy")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Blender K-7"
        table.rows[0].cells[1].text = "14 months"
        path = self.root / "policy.docx"
        doc.save(path)
        text = parse_file(path)
        self.assertIn("Warranty policy", text)
        self.assertIn("Blender K-7 | 14 months", text)

    def test_dataset_prep_never_reads_binary_documents_as_lines(self):
        pdf = self.root / "scan.pdf"
        pdf.write_bytes(_blank_pdf())
        self.assertEqual(_load_records_from_file(pdf), [])
        csv_path = self.root / "excel.csv"
        csv_path.write_bytes(CP1252_CSV)
        records = _load_records_from_file(csv_path)
        self.assertEqual(records[0]["answer"], "Five € per cup")


class UploadCleanFlowTests(unittest.TestCase):
    def setUp(self):
        resp = client.post("/api/projects", json={"name": f"intake-{uuid4().hex[:6]}", "description": ""})
        self.assertEqual(resp.status_code, 201, resp.text)
        self.pid = int(resp.json()["id"])

    def _upload(self, name: str, data: bytes) -> dict:
        resp = client.post(
            f"/api/projects/{self.pid}/ingestion/upload",
            files={"file": (name, data, "application/octet-stream")},
        )
        self.assertIn(resp.status_code, (200, 201), resp.text)
        return resp.json()

    def _process(self, doc_id: int) -> dict:
        resp = client.post(f"/api/projects/{self.pid}/ingestion/documents/{doc_id}/process")
        self.assertEqual(resp.status_code, 200, resp.text)
        return resp.json()

    def test_labelled_csv_keeps_its_columns_through_cleaning(self):
        csv_bytes = (
            "question;answer\n"
            "What does the café charge?;Five € per cup\n"
            "How do I reach support?;Email help@example.com\n"
            "What does the café charge?;Five € per cup\n"
        ).encode("cp1252")
        doc = self._upload("faq.csv", csv_bytes)
        processed = self._process(doc["id"])
        self.assertEqual(processed["status"], "accepted")

        resp = client.post(f"/api/projects/{self.pid}/cleaning/clean", json={"document_id": doc["id"]})
        self.assertEqual(resp.status_code, 200, resp.text)
        body = resp.json()
        self.assertTrue(body.get("structured"))
        self.assertEqual(body["rows_kept"], 2)
        self.assertEqual(body["duplicate_rows_dropped"], 1)

        cleaned_path = Path(settings.DATA_DIR) / "projects" / str(self.pid) / "cleaned" / "cleaned.jsonl"
        rows = [json.loads(line) for line in cleaned_path.read_text(encoding="utf-8").splitlines()]
        self.assertEqual(rows[0]["question"], "What does the café charge?")
        self.assertEqual(rows[0]["answer"], "Five € per cup")
        self.assertNotIn("help@example.com", rows[1]["answer"])  # PII redacted per field
        self.assertEqual(rows[1]["source_document_id"], doc["id"])

    def test_scanned_pdf_is_an_error_not_training_text(self):
        doc = self._upload("scan.pdf", _blank_pdf())
        processed = self._process(doc["id"])
        self.assertEqual(processed["status"], "error")
        self.assertIn("scanned", json.dumps(processed))

    def test_xlsx_upload_is_accepted_and_previewable(self):
        import pandas as pd

        buf = io.BytesIO()
        pd.DataFrame({"text": ["great", "awful"], "label": ["pos", "neg"]}).to_excel(buf, index=False)
        doc = self._upload("reviews.xlsx", buf.getvalue())
        self.assertEqual(self._process(doc["id"])["status"], "accepted")
        sample = client.get(f"/api/projects/{self.pid}/ingestion/documents/{doc['id']}/sample")
        self.assertEqual(sample.status_code, 200, sample.text)
        self.assertEqual({r["label"] for r in sample.json()["rows"]}, {"pos", "neg"})


class ImportWizardUploadTests(unittest.TestCase):
    def setUp(self):
        resp = client.post("/api/projects", json={"name": f"wiz-{uuid4().hex[:6]}", "description": ""})
        self.pid = int(resp.json()["id"])

    def test_upload_stages_file_and_introspects(self):
        resp = client.post(
            f"/api/projects/{self.pid}/dataset-import/upload",
            files={"file": ("faq.csv", CP1252_CSV, "text/csv")},
        )
        self.assertEqual(resp.status_code, 201, resp.text)
        body = resp.json()
        self.assertTrue(body["locator"].startswith("file:"))
        staged = Path(body["locator"][len("file:"):])
        self.assertTrue(staged.exists())
        self.assertIn(f"projects/{self.pid}/imports", staged.as_posix())
        self.assertEqual(body["introspection"]["columns"], ["question", "answer"])

        preview = client.post(
            f"/api/projects/{self.pid}/dataset-import/preview",
            json={
                "locator": body["locator"],
                "mapper_id": "qa_pair_passthrough",
                "field_map": {"question_field": "question", "answer_field": "answer"},
            },
        )
        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertEqual(preview.json()["accepted_count"], 2)

    def test_upload_rejects_documents_with_a_pointer(self):
        resp = client.post(
            f"/api/projects/{self.pid}/dataset-import/upload",
            files={"file": ("manual.pdf", b"%PDF-1.4", "application/pdf")},
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("Upload documents", resp.text)


if __name__ == "__main__":
    unittest.main()
