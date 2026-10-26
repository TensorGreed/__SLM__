"""File parsing utilities for various document types.

Parsers raise :class:`DocumentParseError` when a file can't be turned into
text. They used to return strings like ``"[PDF parse error: …]"``, which the
ingestion step then marked ACCEPTED and cleaning fed to training as if they
were document content.
"""

import hashlib
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from app.services.code_chunker import LANGUAGE_BY_EXTENSION
from app.utils.tabular_io import load_rows, read_text_file, render_row_text


class DocumentParseError(ValueError):
    """The file couldn't be turned into usable text (message is user-facing)."""


def parse_text(file_path: Path) -> str:
    """Read a plain-text file in whatever encoding it was saved with."""
    return read_text_file(file_path)


def parse_markdown(file_path: Path) -> str:
    """Read Markdown file (treated as plain text)."""
    return parse_text(file_path)


def parse_tabular(file_path: Path) -> str:
    """Readable text for CSV/TSV/JSON(L)/XLSX/Parquet — used for previews and
    the ``.extracted.txt`` sidecar. Training keeps the rows themselves
    (cleaning's row-preserving path), not this rendering."""
    try:
        rows = load_rows(file_path)
    except ValueError as exc:
        raise DocumentParseError(str(exc)) from exc
    if not rows:
        raise DocumentParseError(f"{file_path.name} has no data rows.")
    return "\n\n".join(text for text in (render_row_text(row) for row in rows) if text)


def parse_csv(file_path: Path) -> str:
    return parse_tabular(file_path)


# ── OCR for scanned pages ────────────────────────────────────────────────
#
# A scanned PDF has no text layer. When the system has poppler's pdftoppm
# and tesseract (apt: poppler-utils tesseract-ocr), pages without text are
# rendered and OCR'd; otherwise the parse error says what to install.
# BREWSLM_OCR=0 turns it off; BREWSLM_OCR_LANG (default "eng", tesseract
# language codes joined by "+") and BREWSLM_OCR_DPI (default 300) tune it.

OCR_TIMEOUT_SECONDS = 180


def ocr_available() -> bool:
    import os
    import shutil

    if os.environ.get("BREWSLM_OCR", "1").strip().lower() in {"0", "false", "no", "off"}:
        return False
    return bool(shutil.which("pdftoppm") and shutil.which("tesseract"))


def ocr_pdf_pages(file_path: Path, pages: list[int]) -> dict[int, str]:
    """OCR the given 1-based ``pages`` of a PDF: pdftoppm renders each page
    to PNG, tesseract reads it. Returns ``{page: text}``."""
    import os
    import subprocess
    import tempfile

    dpi = str(int(os.environ.get("BREWSLM_OCR_DPI") or 300))
    lang = os.environ.get("BREWSLM_OCR_LANG") or "eng"
    out: dict[int, str] = {}
    with tempfile.TemporaryDirectory(prefix="brewslm-ocr-") as tmp:
        for page in pages:
            prefix = os.path.join(tmp, f"p{page}")
            subprocess.run(
                ["pdftoppm", "-r", dpi, "-f", str(page), "-l", str(page), "-png", "-singlefile", str(file_path), prefix],
                check=True, capture_output=True, timeout=OCR_TIMEOUT_SECONDS,
            )
            result = subprocess.run(
                ["tesseract", f"{prefix}.png", "stdout", "-l", lang, "--psm", "3"],
                check=True, capture_output=True, timeout=OCR_TIMEOUT_SECONDS,
            )
            out[page] = result.stdout.decode("utf-8", errors="replace").strip()
    return out


PAGE_SEPARATOR = "\n\f\n"
"""How :func:`parse_pdf` joins pages: a form feed on its own line, so
cleaning can find page headers / footers per page
(``services/page_furniture``). Cleaning turns it into a blank line — the
join used before — so the cleaned text reads the same."""


def parse_pdf(file_path: Path) -> str:
    """Extract text from a PDF with PyPDF2; pages without a text layer
    (scanned) are OCR'd when tesseract + pdftoppm are installed. Pages are
    joined with :data:`PAGE_SEPARATOR`."""
    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:
        raise DocumentParseError("PDF parsing is unavailable — install PyPDF2 on the server.") from exc
    try:
        reader = PdfReader(str(file_path))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"Couldn't read this PDF ({exc}). Is it password-protected or damaged?") from exc
    empty = [index + 1 for index, text in enumerate(pages) if not text]
    if empty and ocr_available():
        try:
            ocr = ocr_pdf_pages(file_path, empty)
        except Exception as exc:  # noqa: BLE001
            if len(empty) == len(pages):
                raise DocumentParseError(f"This PDF looks scanned and OCR failed ({exc}).") from exc
            ocr = {}
        for page, text in ocr.items():
            pages[page - 1] = text
    text = PAGE_SEPARATOR.join(p for p in pages if p)
    if not text.strip():
        raise DocumentParseError(
            "This PDF has no extractable text — it looks like a scanned/image PDF. "
            + (
                "OCR found no text on its pages either."
                if empty and ocr_available()
                else "Install OCR on the server (apt install tesseract-ocr poppler-utils) to read scans, "
                "or export it as text or DOCX and upload that."
            )
        )
    return text


def parse_docx(file_path: Path) -> str:
    """Extract paragraphs and table cells from DOCX using python-docx."""
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentParseError("DOCX parsing is unavailable — install python-docx on the server.") from exc
    try:
        doc = Document(str(file_path))
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"Couldn't read this Word document ({exc}).") from exc
    blocks = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    text = "\n\n".join(blocks)
    if not text.strip():
        raise DocumentParseError("This Word document has no text.")
    return text


_VOID_TAGS = frozenset({
    "area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr",
})
_NAV_TAGS = frozenset({"nav", "header", "footer", "aside"})
_NAV_ROLES = frozenset({"navigation", "banner", "contentinfo", "search"})
_NAV_WORDS = frozenset({"breadcrumb", "breadcrumbs", "nav", "navbar", "navigation", "menu", "sidebar", "toc", "footer"})
# Never navigation, whatever their class / role says: the page itself and its content landmarks.
_NEVER_NAV_TAGS = frozenset({"html", "body", "main", "article"})
# Only these containers are matched by class / id words — a ``<p class="nav-hint">`` or
# ``<a class="menu-link">`` is content-sized and ``div.layout.with-sidebar`` is handled by the fallback.
_NAV_CLASS_TAGS = frozenset({"div", "ul", "ol", "section", "span", "table", "form"})
_SKIP_TAGS = frozenset({"script", "style", "noscript", "template", "svg", "head"})
_BLOCK_TAGS = frozenset({
    "p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "pre", "td", "th",
})
# Start tags that implicitly close an open ``<p>``.
_P_CLOSERS = frozenset({
    "address", "article", "aside", "blockquote", "details", "div", "dl", "fieldset", "figure", "footer", "form",
    "h1", "h2", "h3", "h4", "h5", "h6", "header", "hr", "main", "nav", "ol", "p", "pre", "section", "table", "ul", "li",
})
_SCOPE_BOUNDARY = frozenset({"table", "td", "th", "caption", "html", "body", "button", "template"})
# start tag -> (open tags it implicitly closes, boundaries the search won't cross)
_IMPLIED_CLOSE = {
    "li": (frozenset({"li"}), frozenset({"ul", "ol"})),
    "dt": (frozenset({"dt", "dd"}), frozenset({"dl"})),
    "dd": (frozenset({"dt", "dd"}), frozenset({"dl"})),
    "tr": (frozenset({"tr", "td", "th"}), frozenset({"table"})),
    "td": (frozenset({"td", "th"}), frozenset({"tr", "table"})),
    "th": (frozenset({"td", "th"}), frozenset({"tr", "table"})),
    "option": (frozenset({"option"}), frozenset({"select"})),
}


def _is_main_element(tag: str, attrs: list) -> bool:
    if tag == "main":
        return True
    return dict((name, value or "") for name, value in attrs).get("role", "").strip().lower() == "main"


def _is_nav_element(tag: str, attrs: list, *, use_class: bool = True) -> bool:
    """Whether an element is site chrome. ``html`` / ``body`` / ``main`` / ``article`` and
    ``role="main"`` are never chrome; class / id words (``sidebar``, ``nav``…) only count on
    container tags and only when ``use_class``."""
    if tag in _NEVER_NAV_TAGS or _is_main_element(tag, attrs):
        return False
    if tag in _NAV_TAGS:
        return True
    values = {name: (value or "") for name, value in attrs}
    if values.get("role", "").strip().lower() in _NAV_ROLES:
        return True
    if not use_class or tag not in _NAV_CLASS_TAGS:
        return False
    for name in ("class", "id"):
        # Token match (``nav-bar``, ``main_menu`` split into words), so ``canvas`` or ``menuhead``-style
        # names that merely contain the letters don't count.
        words = set(re.split(r"[^a-z0-9]+", values.get(name, "").lower()))
        if words & _NAV_WORDS:
            return True
    return False


class _Node:
    __slots__ = ("tag", "attrs", "children")

    def __init__(self, tag: str, attrs: list) -> None:
        self.tag = tag
        self.attrs = attrs
        self.children: list = []  # _Node | str


class _HTMLText(HTMLParser):
    """Builds a light element tree (void tags, stray end tags and implicitly closed
    ``<p>`` / ``<li>`` / table cells handled), so navigation can be dropped by element
    without an unclosed tag swallowing the rest of the page."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root = _Node("#root", [])
        self._stack: list[_Node] = [self.root]

    def _pop_to(self, targets: frozenset, boundaries: frozenset) -> None:
        for index in range(len(self._stack) - 1, 0, -1):
            tag = self._stack[index].tag
            if tag in targets:
                del self._stack[index:]
                return
            if tag in boundaries:
                return

    def handle_starttag(self, tag, attrs):  # noqa: ANN001
        if tag in _P_CLOSERS:
            self._pop_to(frozenset({"p"}), _SCOPE_BOUNDARY)
        if tag in _IMPLIED_CLOSE:
            targets, boundaries = _IMPLIED_CLOSE[tag]
            self._pop_to(targets, boundaries)
        node = _Node(tag, attrs)
        self._stack[-1].children.append(node)
        if tag not in _VOID_TAGS:
            self._stack.append(node)

    def handle_startendtag(self, tag, attrs):  # noqa: ANN001
        self._stack[-1].children.append(_Node(tag, attrs))

    def handle_endtag(self, tag):  # noqa: ANN001
        for index in range(len(self._stack) - 1, 0, -1):
            if self._stack[index].tag == tag:
                del self._stack[index:]
                return

    def handle_data(self, data):  # noqa: ANN001
        self._stack[-1].children.append(data)


def _find_main(root: _Node) -> list[_Node]:
    """Outermost ``<main>`` / ``role="main"`` elements, else outermost ``<article>``s."""
    for match in (lambda n: _is_main_element(n.tag, n.attrs), lambda n: n.tag == "article"):
        found: list[_Node] = []
        stack = [root]
        while stack:
            node = stack.pop()
            for child in reversed(node.children):
                if not isinstance(child, _Node):
                    continue
                if match(child):
                    found.append(child)
                else:
                    stack.append(child)
        if found:
            # stack-based walk visits out of document order; restore it
            order = {id(n): i for i, n in enumerate(_walk(root))}
            return sorted(found, key=lambda n: order[id(n)])
    return []


def _walk(root: _Node):
    stack = [root]
    while stack:
        node = stack.pop()
        yield node
        stack.extend(c for c in reversed(node.children) if isinstance(c, _Node))


def _render(roots: list[_Node], *, mode: int) -> str:
    """Visible text under ``roots``. ``mode``: 0 drops navigation by tag, role and
    class / id; 1 by tag and role only; 2 drops nothing."""
    parts: list[str] = []
    for root in roots:
        stack: list = [root]
        while stack:
            item = stack.pop()
            if item is None:
                parts.append("\n")  # closing a block element
                continue
            if isinstance(item, str):
                parts.append(item)
                continue
            tag = item.tag
            if tag in _SKIP_TAGS:
                continue
            if mode < 2 and item is not root and _is_nav_element(tag, item.attrs, use_class=mode == 0):
                continue
            if tag in _BLOCK_TAGS:
                parts.append("\n")
                if tag not in _VOID_TAGS:
                    stack.append(None)
            stack.extend(reversed(item.children))
    lines = [" ".join(line.split()) for line in "".join(parts).splitlines()]
    return "\n".join(line for line in lines if line)


def parse_html(file_path: Path, *, strip_navigation: bool = False) -> str:
    """Visible text of an HTML page (scripts/styles dropped).

    ``strip_navigation`` (crawled pages — a site's chrome repeats on every
    page and would dominate retrieval) also drops ``nav`` / ``header`` /
    ``footer`` / ``aside``, elements with a navigation ``role`` (navigation,
    banner, contentinfo, search) and container elements whose class / id names
    a breadcrumb, nav, menu, sidebar, toc or footer; when the page has a
    ``<main>`` / ``role="main"`` (else ``<article>``) only that text is kept.
    ``html`` / ``body`` / ``main`` / ``article`` are never dropped. When
    stripping would leave (almost) nothing — a ``layout with-sidebar`` wrapper
    around the whole page — it retries by tag / role only, then keeps everything.
    Uploads keep the default (everything visible)."""
    return html_to_text(read_text_file(file_path), strip_navigation=strip_navigation)


_MIN_STRIPPED_FRACTION = 0.2


def html_to_text(markup: str, *, strip_navigation: bool = False) -> str:
    """:func:`parse_html` on markup already in memory."""
    extractor = _HTMLText()
    extractor.feed(markup)
    extractor.close()
    root = extractor.root
    text = _render([root], mode=2)
    if strip_navigation and text:
        mains = _find_main(root)
        roots = mains or [root]
        for mode in (0, 1):
            stripped = _render(roots, mode=mode)
            # With a <main> the chrome is already gone: any text is enough. Without one,
            # stripping must not have eaten (nearly) the whole page.
            enough = bool(stripped) if mains else len(stripped) >= _MIN_STRIPPED_FRACTION * len(text)
            if enough:
                text = stripped
                break
    if not text.strip():
        raise DocumentParseError("This HTML page has no visible text.")
    return text


def parse_code(file_path: Path) -> str:
    """Source / config text exactly as written — no whitespace collapsing, so
    indentation and line numbers survive (the code chunker cites them)."""
    text = read_text_file(file_path)
    if not text.strip():
        raise DocumentParseError(f"{file_path.name} is empty.")
    return text


# File extension → parser mapping
PARSERS: dict[str, Any] = {
    ".txt": parse_text,
    ".md": parse_markdown,
    ".markdown": parse_markdown,
    ".rst": parse_text,
    ".csv": parse_tabular,
    ".tsv": parse_tabular,
    ".json": parse_tabular,
    ".jsonl": parse_tabular,
    ".xlsx": parse_tabular,
    ".parquet": parse_tabular,
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".html": parse_html,
    ".htm": parse_html,
}
# Source / config files (``.json`` stays a row file for uploads; a repo import
# marks its documents ``content_kind="code"`` and parses those as code).
for _ext, _language in LANGUAGE_BY_EXTENSION.items():
    if _ext != ".json":
        PARSERS.setdefault(_ext, parse_code)

SUPPORTED_EXTENSIONS = set(PARSERS.keys())


def parse_file(file_path: Path, **options: Any) -> str:
    """Parse a file based on its extension. Returns extracted text; raises
    :class:`DocumentParseError` when there's no usable text. ``options`` go to
    the parser (``strip_navigation`` for HTML); ``content_kind="code"`` parses
    any file as code text (a repo's ``.json`` / ``.yaml`` are not row files)."""
    ext = file_path.suffix.lower()
    if options.pop("content_kind", None) == "code":
        return parse_code(file_path)
    parser = PARSERS.get(ext)
    if not parser:
        raise ValueError(f"Unsupported file type: {ext}")
    if parser is parse_html:
        return parser(file_path, **{k: v for k, v in options.items() if k == "strip_navigation"})
    return parser(file_path)


def compute_file_hash(file_path: Path) -> str:
    """Compute SHA-256 hash of file contents."""
    sha256 = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            sha256.update(chunk)
    return sha256.hexdigest()


def get_file_type(filename: str) -> str:
    """Return the file extension without dot."""
    return Path(filename).suffix.lower().lstrip(".")
